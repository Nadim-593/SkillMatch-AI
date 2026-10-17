import io
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.document_parser import extract_document_text
from app.extraction import _safe_parse_json, _heuristic_skill_extract
from app.normalization import normalize_skill

client = TestClient(app)

def test_taxonomy_normalization():
    assert normalize_skill("js") == "JavaScript"
    assert normalize_skill("JS") == "JavaScript"
    assert normalize_skill("react.js") == "React"
    assert normalize_skill("nodejs") == "Node.js"
    assert normalize_skill("postgres") == "PostgreSQL"
    assert normalize_skill("k8s") == "Kubernetes"
    assert normalize_skill("unknown_custom_skill") == "unknown_custom_skill"

def test_heuristic_skill_extract():
    sample_text = (
        "Experienced software developer with 4 years in Python, React, and PostgreSQL. "
        "Built microservices with Docker, Kubernetes, and AWS."
    )
    extracted = _heuristic_skill_extract(sample_text)
    names = {s["normalized_name"] for s in extracted}
    assert "Python" in names
    assert "React" in names
    assert "PostgreSQL" in names
    assert "Docker" in names
    assert "Kubernetes" in names
    assert "AWS" in names
    # Evidence must be present
    for s in extracted:
        assert len(s["evidence"]) > 0

def test_safe_parse_json_markdown_blocks():
    raw_markdown = """```json
[
  {"skill_name": "react", "normalized_name": "React", "category": "Frontend", "evidence": "React developer", "confidence": 0.9}
]
```"""
    parsed = _safe_parse_json(raw_markdown)
    assert isinstance(parsed, list)
    assert len(parsed) == 1
    assert parsed[0]["normalized_name"] == "React"

def test_safe_parse_json_embedded_json():
    raw_text = """Here is the result of your skill extraction:
[
  {"skill_name": "Python", "normalized_name": "Python", "category": "Programming Language", "evidence": "Python 3", "confidence": 0.95}
]
Hope this helps!"""
    parsed = _safe_parse_json(raw_text)
    assert isinstance(parsed, list)
    assert parsed[0]["normalized_name"] == "Python"

def test_safe_parse_json_invalid():
    with pytest.raises(ValueError):
        _safe_parse_json("This is pure prose without any JSON formatting whatsoever.")

def test_extract_cv_api_valid():
    cv_text = (
        "Full stack engineer. Proficient in Node.js, Express.js, TypeScript, and MongoDB. "
        "Implemented CI/CD pipelines with GitHub Actions."
    )
    res = client.post("/api/v1/extract/cv", json={"text": cv_text})
    assert res.status_code == 200
    data = res.json()
    assert data["success"] is True
    assert len(data["skills"]) > 0
    names = [s["normalized_name"] for s in data["skills"]]
    assert "Node.js" in names or "TypeScript" in names

def test_extract_cv_api_empty_text():
    res = client.post("/api/v1/extract/cv", json={"text": "   "})
    assert res.status_code == 422 or res.status_code == 400

def test_extract_cv_api_too_short():
    res = client.post("/api/v1/extract/cv", json={"text": "short"})
    assert res.status_code == 422 or res.status_code == 400

def test_pdf_document_parsing():
    import fitz
    # Create valid PDF in memory
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "John Doe — Senior Backend Developer with 5 years in Python and Docker.")
    pdf_bytes = doc.tobytes()
    doc.close()

    text = extract_document_text(pdf_bytes, "resume.pdf")
    assert "Senior Backend Developer" in text
    assert "Python" in text
    assert "Docker" in text

def test_docx_document_parsing():
    import docx
    # Create valid DOCX in memory
    doc = docx.Document()
    doc.add_heading("Jane Smith — Full Stack Engineer", level=1)
    doc.add_paragraph("Skills: React, JavaScript, Node.js, and PostgreSQL database management.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    text = extract_document_text(docx_bytes, "resume.docx")
    assert "Full Stack Engineer" in text
    assert "React" in text
    assert "PostgreSQL" in text

def test_corrupted_document_handling():
    with pytest.raises(ValueError, match="Invalid or corrupted"):
        extract_document_text(b"not a valid pdf binary content", "corrupted.pdf")

    with pytest.raises(ValueError, match="Invalid or corrupted"):
        extract_document_text(b"not a valid docx binary content", "corrupted.docx")

def test_empty_document_handling():
    with pytest.raises(ValueError, match="empty"):
        extract_document_text(b"", "empty.pdf")

def test_unsupported_file_extension():
    with pytest.raises(ValueError, match="Unsupported file format"):
        extract_document_text(b"some content", "resume.exe")
