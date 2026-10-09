import io
import logging
from typing import Tuple

logger = logging.getLogger(__name__)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    try:
        import fitz
    except ImportError:
        raise RuntimeError("PyMuPDF (fitz) is not installed in the environment.")
    
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        if doc.page_count == 0:
            raise ValueError("PDF document has 0 pages.")
        
        text_parts = []
        for page_idx in range(doc.page_count):
            page = doc.load_page(page_idx)
            text = page.get_text()
            if text:
                text_parts.append(text.strip())
        doc.close()
        return "\n\n".join(text_parts).strip()
    except Exception as e:
        logger.error(f"Failed to parse PDF: {e}")
        raise ValueError(f"Invalid or corrupted PDF file: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    try:
        import docx
    except ImportError:
        raise RuntimeError("python-docx is not installed in the environment.")
    
    try:
        file_stream = io.BytesIO(file_bytes)
        doc = docx.Document(file_stream)
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        return "\n\n".join(text_parts).strip()
    except Exception as e:
        logger.error(f"Failed to parse DOCX: {e}")
        raise ValueError(f"Invalid or corrupted DOCX file: {str(e)}")

def extract_document_text(file_bytes: bytes, filename: str) -> str:
    """Parse text from either PDF or DOCX format."""
    if not file_bytes or len(file_bytes) == 0:
        raise ValueError("Uploaded document is completely empty (0 bytes).")
    
    fname_lower = filename.lower()
    if fname_lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif fname_lower.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise ValueError("Unsupported file format. Only PDF and DOCX files are allowed.")
    
    clean_text = text.strip()
    if len(clean_text) < 15:
        raise ValueError("Document contains insufficient readable text or is an empty/scanned document.")
    
    return clean_text
